import re
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import and_, delete, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.deletion_log import record_deletion
from app.routers.auth import require_admin
from app.unregistered_number import assign_missing_unregistered_numbers, hold_unregistered_number
from app.models import Company, Employee, JoinChild, JoinEvent
from app.party import company_on_clause, in_department, labels
from app.pass_pdf import build_passes_pdf
from app.schemas import (
    DepartmentOut,
    JoinEventDeleteResult,
    JoinEventUploadResult,
    PassAllPrintRequest,
    PassOut,
    PassPrintRequest,
)

router = APIRouter(prefix="/api/passes", tags=["passes"])


def _department_rows(company_code: str, department: str, db: Session, *, printed: bool):
    print_status = JoinEvent.print_status == "Printed" if printed else JoinEvent.print_status != "Printed"
    statement = (
        select(JoinEvent, Employee, Company)
        .outerjoin(Employee, JoinEvent.employee_id == Employee.employee_id)
        .outerjoin(Company, company_on_clause())
        .where(
            in_department(company_code, department),
            JoinEvent.approve_status == "Approved",
            print_status,
        )
    )
    rows = list(db.execute(statement))
    rows.sort(key=lambda row: _pass_sort_key(*row))
    return rows


def _approved_rows(company_code: str, department: str, db: Session):
    return _department_rows(company_code, department, db, printed=False)


def _printed_rows(company_code: str, department: str, db: Session):
    return _department_rows(company_code, department, db, printed=True)


def _employee_id_sort(employee_id: str):
    value = (employee_id or "").strip()
    if value.isdigit():
        return (0, int(value))
    if value:
        return (1, value.casefold())
    return (2, "")


def _pass_sort_key(event: JoinEvent, employee: Employee | None, company: Company | None):
    employee_id, employee_name, department, company_code, _company_name = labels(event, employee, company)
    relation = (event.relation or "").strip().casefold()
    return (
        (company_code or "").casefold(),
        (department or "").casefold(),
        _employee_id_sort(employee_id),
        (employee_name or "").casefold(),
        0 if relation == "employee" else 1,
        (event.name or "").casefold(),
        event.unique_event_id,
    )


def _all_department_rows(company_code: str, db: Session):
    statement = (
        select(JoinEvent, Employee, Company)
        .outerjoin(Employee, JoinEvent.employee_id == Employee.employee_id)
        .outerjoin(Company, company_on_clause())
        .where(
            JoinEvent.approve_status == "Approved",
            JoinEvent.print_status != "Printed",
        )
    )
    if company_code:
        statement = statement.where(
            or_(
                Employee.company_code == company_code,
                and_(JoinEvent.employee_id.is_(None), JoinEvent.company == company_code),
            )
        )
    rows = []
    for event, employee, company in db.execute(statement):
        _employee_id, _employee_name, _department, company_code, _company_name = labels(event, employee, company)
        if company_code and company_code.strip():
            rows.append((event, employee, company))
    rows.sort(key=lambda row: _pass_sort_key(*row))
    return rows


def _to_pass(event: JoinEvent, employee: Employee | None, company: Company) -> PassOut:
    employee_id, employee_name, department, company_code, company_name = labels(event, employee, company)
    return PassOut(
        unique_event_id=event.unique_event_id,
        employee_id=employee_id,
        employee_name=employee_name,
        department=department,
        company_code=company_code,
        company_name=company_name,
        name=event.name,
        relation=event.relation,
        approve_status=event.approve_status,
        print_status=event.print_status,
        unregistered_number=event.unregistered_number or "",
    )


@router.get("/pending-departments", response_model=list[DepartmentOut])
def list_pending_print_departments(company_code: str = "", db: Session = Depends(get_db)):
    company_code = company_code.strip()
    registered = (
        select(Employee.company_code, Employee.department)
        .join(JoinEvent, JoinEvent.employee_id == Employee.employee_id)
        .where(
            JoinEvent.approve_status == "Approved",
            JoinEvent.print_status != "Printed",
        )
    )
    unregistered = select(JoinEvent.company, JoinEvent.department).where(
        JoinEvent.employee_id.is_(None),
        JoinEvent.approve_status == "Approved",
        JoinEvent.print_status != "Printed",
    )
    if company_code:
        registered = registered.where(Employee.company_code == company_code)
        unregistered = unregistered.where(JoinEvent.company == company_code)
    pairs = {
        (company.strip(), department.strip())
        for company, department in db.execute(registered)
        if company and str(company).strip() and department and str(department).strip()
    }
    pairs.update(
        (str(company).strip(), str(department).strip())
        for company, department in db.execute(unregistered)
        if company and str(company).strip() and department and str(department).strip()
    )
    return [
        DepartmentOut(department=department, company_code=company)
        for company, department in sorted(pairs, key=lambda item: (item[0].casefold(), item[1].casefold()))
    ]


@router.get("/all-departments", response_model=list[PassOut])
def list_all_department_passes(company_code: str = "", db: Session = Depends(get_db)):
    return [
        _to_pass(event, employee, company)
        for event, employee, company in _all_department_rows(company_code.strip(), db)
    ]


@router.post("/all-departments/pdf")
def print_all_department_passes(payload: PassAllPrintRequest, db: Session = Depends(get_db)):
    wanted = set(payload.unique_event_ids)
    chosen = [
        (event, employee, company)
        for event, employee, company in _all_department_rows(payload.company_code, db)
        if event.unique_event_id in wanted
    ]
    found = {event.unique_event_id for event, _, _ in chosen}
    if found != wanted:
        raise HTTPException(
            status_code=400,
            detail="Select approved entries that are still waiting to print",
        )
    pdf = build_passes_pdf(_pass_cards(chosen))
    for event, _, _ in chosen:
        event.print_status = "Printed"
    db.commit()
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="event-passes.pdf"'},
    )


@router.get("", response_model=list[PassOut])
def list_approved_passes(
    company_code: str = Query(min_length=1),
    department: str = Query(min_length=1),
    db: Session = Depends(get_db),
):
    return [
        _to_pass(event, employee, company)
        for event, employee, company in _approved_rows(company_code.strip(), department.strip(), db)
        if company is not None
    ]


@router.post("/pdf")
def print_passes(payload: PassPrintRequest, db: Session = Depends(get_db)):
    wanted = set(payload.unique_event_ids)
    chosen = []
    for event, employee, company in _approved_rows(payload.company_code, payload.department, db):
        if event.unique_event_id in wanted and company is not None:
            chosen.append((event, employee, company))
    found = {event.unique_event_id for event, _, _ in chosen}
    missing = wanted - found
    if missing:
        raise HTTPException(
            status_code=400,
            detail="Select approved entries from this department only",
        )

    cards = []
    for event, employee, company in chosen:
        employee_id, employee_name, department, _company_code, _company_name = labels(event, employee, company)
        cards.append(
            {
                "unique_event_id": event.unique_event_id,
                "department": department,
                "employee_name": employee_name,
                "employee_id": employee_id,
                "attendee_name": event.name,
                "relation": event.relation,
                "category": event.category,
                "admission_group": event.admission_group,
                "event": event.event,
                "unregistered_number": event.unregistered_number or "",
            }
        )
    pdf = build_passes_pdf(cards)
    for event, _, _ in chosen:
        event.print_status = "Printed"
    db.commit()
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="event-passes.pdf"'},
    )


def _pass_cards(chosen):
    cards = []
    for event, employee, _company in chosen:
        employee_id, employee_name, department, _company_code, _company_name = labels(event, employee, _company)
        cards.append(
            {
                "unique_event_id": event.unique_event_id,
                "department": department,
                "employee_name": employee_name,
                "employee_id": employee_id,
                "attendee_name": event.name,
                "relation": event.relation,
                "category": event.category,
                "admission_group": event.admission_group,
                "event": event.event,
                "unregistered_number": event.unregistered_number or "",
            }
        )
    return cards


@router.get("/reprints", response_model=list[PassOut])
def list_printed_passes(
    company_code: str = Query(min_length=1),
    department: str = Query(min_length=1),
    db: Session = Depends(get_db),
):
    return [
        _to_pass(event, employee, company)
        for event, employee, company in _printed_rows(company_code.strip(), department.strip(), db)
        if company is not None
    ]


@router.post("/reprint-pdf")
def reprint_passes(payload: PassPrintRequest, db: Session = Depends(get_db)):
    wanted = set(payload.unique_event_ids)
    chosen = []
    for event, employee, company in _printed_rows(payload.company_code, payload.department, db):
        if event.unique_event_id in wanted and company is not None:
            chosen.append((event, employee, company))
    found = {event.unique_event_id for event, _, _ in chosen}
    missing = wanted - found
    if missing:
        raise HTTPException(
            status_code=400,
            detail="Select printed entries from this department only",
        )
    pdf = build_passes_pdf(_pass_cards(chosen))
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="event-passes.pdf"'},
    )


JOIN_COLUMNS = (
    ("UniqueEventId", "unique_event_id"),
    ("EmployeeId", "employee_id"),
    ("Name", "name"),
    ("Relation", "relation"),
    ("Department", "department"),
    ("EmployeeName", "employee_name"),
    ("Company", "company"),
    ("Category", "category"),
    ("AdmissionGroup", "admission_group"),
    ("ApproveStatus", "approve_status"),
    ("PrintStatus", "print_status"),
    ("JoinStatus", "join_status"),
    ("MealsStatus", "meals_status"),
    ("ReportIssue", "report_issue"),
    ("Event", "event"),
    ("EventYear", "event_year"),
    ("UnregisteredNumber", "unregistered_number"),
)
JOIN_REQUIRED = ("name", "relation")
JOIN_LIMITS = {
    "employee_id": 50,
    "name": 200,
    "relation": 20,
    "department": 100,
    "employee_name": 200,
    "company": 50,
    "category": 20,
    "admission_group": 20,
    "approve_status": 20,
    "print_status": 20,
    "join_status": 20,
    "meals_status": 20,
    "report_issue": 50,
    "event": 100,
    "unregistered_number": 20,
}
JOIN_DEFAULTS = {
    "department": "",
    "employee_name": "",
    "company": "",
    "category": "Registered",
    "admission_group": "Employee",
    "approve_status": "Pending",
    "print_status": "Pending",
    "join_status": "Pending",
    "meals_status": "Pending",
    "report_issue": "",
    "event": "Annual Day 2026",
    "event_year": 2026,
    "unregistered_number": "",
}
JOIN_MAX_BYTES = 5 * 1024 * 1024
JOIN_SHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _header_key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _reset_join_sequence() -> None:
    from app.database import engine

    with engine.begin() as connection:
        connection.execute(text("ALTER TABLE `Joinevents` AUTO_INCREMENT = 1000000000"))


def _join_values(event: JoinEvent) -> list:
    values = []
    for _label, field in JOIN_COLUMNS:
        value = getattr(event, field)
        if field == "employee_id" and value is None:
            value = ""
        values.append(value)
    return values


@router.get("/join-events")
def download_join_events(_admin: str = Depends(require_admin), db: Session = Depends(get_db)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Joinevents"
    sheet.append([label for label, _field in JOIN_COLUMNS])
    for event in db.scalars(select(JoinEvent).order_by(JoinEvent.unique_event_id)).all():
        sheet.append(_join_values(event))
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type=JOIN_SHEET,
        headers={"Content-Disposition": 'attachment; filename="joinevents.xlsx"'},
    )


@router.delete("/join-events", response_model=JoinEventDeleteResult)
def delete_all_join_events(
    request: Request,
    _admin: str = Depends(require_admin),
    db: Session = Depends(get_db),
):
    deleted = db.execute(delete(JoinEvent)).rowcount or 0
    children_deleted = db.execute(delete(JoinChild)).rowcount or 0
    record_deletion(
        db,
        request,
        "Delete JoinEvents",
        f"Deleted {deleted} join-event rows and {children_deleted} child rows.",
    )
    db.commit()
    _reset_join_sequence()
    return JoinEventDeleteResult(deleted=deleted, children_deleted=children_deleted)


@router.post("/join-events", response_model=JoinEventUploadResult)
async def upload_join_events(
    file: UploadFile = File(...),
    _admin: str = Depends(require_admin),
    db: Session = Depends(get_db),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Choose an Excel .xlsx file")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The file is empty")
    if len(content) > JOIN_MAX_BYTES:
        raise HTTPException(status_code=400, detail="The file is larger than 5 MB")

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that Excel file")

    try:
        sheet = workbook.active
        row_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(row_iter)
        except StopIteration:
            raise HTTPException(status_code=400, detail="The worksheet is empty")

        fields = {field: label for label, field in JOIN_COLUMNS}
        columns = {}
        for index, cell in enumerate(header):
            key = _header_key(cell)
            for field, label in fields.items():
                if key == _header_key(label) and field not in columns:
                    columns[field] = index
        missing = [fields[field] for field in JOIN_REQUIRED if field not in columns]
        if missing:
            raise HTTPException(
                status_code=400,
                detail="The first row must include these columns: " + ", ".join(missing),
            )

        parsed = []
        errors = []
        seen_ids = set()
        for offset, row in enumerate(row_iter, start=2):
            values = {
                field: _cell_text(row[index] if index < len(row) else None)
                for field, index in columns.items()
            }
            if not any(values.values()):
                continue

            row_errors = []
            event_id = None
            if "unique_event_id" in values and values["unique_event_id"]:
                try:
                    event_id = int(values["unique_event_id"])
                except ValueError:
                    row_errors.append("UniqueEventId must be a whole number")
                if event_id is not None and event_id < 1:
                    row_errors.append("UniqueEventId must be greater than zero")
                elif event_id in seen_ids:
                    row_errors.append(f"UniqueEventId {event_id} is repeated")
            if not values.get("name"):
                row_errors.append("Name is required")
            if not values.get("relation"):
                row_errors.append("Relation is required")

            event_year = JOIN_DEFAULTS["event_year"]
            if "event_year" in values and values["event_year"]:
                try:
                    event_year = int(values["event_year"])
                except ValueError:
                    row_errors.append("EventYear must be a whole number")

            cleaned = {"unique_event_id": event_id, "event_year": event_year}
            for field, limit in JOIN_LIMITS.items():
                if field == "employee_id":
                    text_value = values.get(field, "")
                elif field in ("name", "relation"):
                    text_value = values.get(field, "")
                else:
                    text_value = values.get(field) or JOIN_DEFAULTS[field]
                if len(text_value) > limit:
                    row_errors.append(f"{fields[field]} is longer than {limit} characters")
                cleaned[field] = text_value

            if row_errors:
                errors.append(f"Row {offset}: {'; '.join(row_errors)}")
                continue
            if event_id is not None:
                seen_ids.add(event_id)
            parsed.append(cleaned)

        if errors:
            preview = " ".join(errors[:5])
            extra = f" {len(errors) - 5} more." if len(errors) > 5 else ""
            raise HTTPException(status_code=400, detail=preview + extra)
        if not parsed:
            raise HTTPException(status_code=400, detail="The worksheet has no join-event rows")

        employee_ids = {row["employee_id"] for row in parsed if row["employee_id"]}
        known_employees = set()
        if employee_ids:
            known_employees = set(
                db.scalars(select(Employee.employee_id).where(Employee.employee_id.in_(employee_ids))).all()
            )
        unknown = sorted(employee_ids - known_employees)
        if unknown:
            shown = ", ".join(unknown[:5])
            extra = f" and {len(unknown) - 5} more" if len(unknown) > 5 else ""
            raise HTTPException(
                status_code=400,
                detail=f"These EmployeeId values are not in EmployeeMaster: {shown}{extra}",
            )

        event_ids = [row["unique_event_id"] for row in parsed if row["unique_event_id"] is not None]
        existing = {}
        if event_ids:
            for event in db.scalars(select(JoinEvent).where(JoinEvent.unique_event_id.in_(event_ids))).all():
                existing[event.unique_event_id] = event

        inserted = 0
        updated = 0
        for row in parsed:
            event_id = row.pop("unique_event_id")
            employee_id = row.pop("employee_id") or None
            if event_id is not None and event_id in existing:
                event = existing[event_id]
                event.employee_id = employee_id
                for field, value in row.items():
                    setattr(event, field, value)
                updated += 1
            else:
                payload = {"employee_id": employee_id, **row}
                if event_id is not None:
                    payload["unique_event_id"] = event_id
                db.add(JoinEvent(**payload))
                inserted += 1
        try:
            with hold_unregistered_number(db):
                assign_missing_unregistered_numbers(db)
                db.commit()
        except IntegrityError:
            db.rollback()
            raise HTTPException(status_code=409, detail="JoinEvents could not be saved")
        except RuntimeError:
            db.rollback()
            raise HTTPException(status_code=503, detail="Could not assign an unregistered number. Try again.")
        _reset_join_sequence()
        return JoinEventUploadResult(inserted=inserted, updated=updated)
    finally:
        workbook.close()
