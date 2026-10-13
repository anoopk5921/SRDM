import re
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from openpyxl import load_workbook
from sqlalchemy import delete, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.deletion_log import record_deletion
from app.routers.auth import require_admin
from app.models import Company, Employee, JoinChild, JoinEvent
from app.schemas import EmployeeDeleteResult, EmployeeOut, EmployeeStrengthUpdate, EmployeeUploadResult

router = APIRouter(prefix="/api/employees", tags=["employees"])

HEADER_FIELDS = {
    "employeeid": "employee_id",
    "employeename": "employee_name",
    "department": "department",
    "mobile": "mobile",
    "mobileno": "mobile",
    "mobilenumber": "mobile",
    "companycode": "company_code",
    "allowedstrength": "allowed_strength",
}
REQUIRED = list(HEADER_FIELDS.values())
MAX_BYTES = 5 * 1024 * 1024


def _header_key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


@router.get("", response_model=list[EmployeeOut])
def list_employees(q: str | None = None, db: Session = Depends(get_db)):
    statement = select(Employee).order_by(Employee.employee_id)
    if q and q.strip():
        term = f"%{q.strip()}%"
        statement = statement.where(
            or_(
                Employee.employee_id.like(term),
                Employee.employee_name.like(term),
                Employee.department.like(term),
                Employee.company_code.like(term),
            )
        )
    return db.scalars(statement).all()


@router.delete("", response_model=EmployeeDeleteResult)
def delete_all_employees(
    request: Request,
    _admin: str = Depends(require_admin),
    db: Session = Depends(get_db),
):
    try:
        events = db.execute(delete(JoinEvent))
        children = db.execute(delete(JoinChild))
        employees = db.execute(delete(Employee))
        record_deletion(
            db,
            request,
            "Delete employees",
            (
                f"Deleted {employees.rowcount or 0} employees, "
                f"{events.rowcount or 0} join-event rows, "
                f"and {children.rowcount or 0} child rows."
            ),
        )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Employees could not be deleted because other records still refer to them",
        )
    return EmployeeDeleteResult(
        deleted_employees=employees.rowcount or 0,
        deleted_events=events.rowcount or 0,
    )


@router.put("/strength", response_model=EmployeeOut)
def update_allowed_strength(payload: EmployeeStrengthUpdate, db: Session = Depends(get_db)):
    employee = db.get(Employee, payload.employee_id)
    if employee is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    employee.allowed_strength = payload.allowed_strength
    db.commit()
    db.refresh(employee)
    return employee


@router.get("/{employee_id}", response_model=EmployeeOut)
def get_employee(employee_id: str, db: Session = Depends(get_db)):
    employee = db.get(Employee, employee_id.strip())
    if employee is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    return employee


@router.post("/upload", response_model=EmployeeUploadResult)
async def upload_employees(file: UploadFile = File(...), db: Session = Depends(get_db)):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Choose an Excel .xlsx file")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The file is empty")
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="The file is larger than 5 MB")

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that Excel file")

    try:
        sheet = workbook.active
        row_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(row_iter)
        except StopIteration:
            raise HTTPException(status_code=400, detail="The worksheet is empty")

        columns = {}
        for index, cell in enumerate(header):
            field = HEADER_FIELDS.get(_header_key(cell))
            if field and field not in columns:
                columns[field] = index
        missing = [name for name in REQUIRED if name not in columns]
        if missing:
            labels = ", ".join(
                {
                    "employee_id": "EmployeeId",
                    "employee_name": "EmployeeName",
                    "department": "Department",
                    "mobile": "Mobile",
                    "company_code": "CompanyCode",
                    "allowed_strength": "AllowedStrength",
                }[name]
                for name in missing
            )
            raise HTTPException(
                status_code=400,
                detail=f"The first row must include these columns: {labels}",
            )

        company_codes = set(db.scalars(select(Company.company_code)).all())
        parsed = []
        errors = []
        seen_ids = set()

        for offset, row in enumerate(row_iter, start=2):
            values = {field: _cell_text(row[index] if index < len(row) else None) for field, index in columns.items()}
            if not any(values.values()):
                continue

            row_errors = []
            if not values["employee_id"]:
                row_errors.append("EmployeeId is required")
            elif values["employee_id"] in seen_ids:
                row_errors.append(f"EmployeeId {values['employee_id']} is repeated")
            if not values["employee_name"]:
                row_errors.append("EmployeeName is required")
            if not values["department"]:
                row_errors.append("Department is required")
            if not values["company_code"]:
                row_errors.append("CompanyCode is required")
            elif values["company_code"] not in company_codes:
                row_errors.append(f"CompanyCode {values['company_code']} is not in the company master")

            strength_text = values["allowed_strength"]
            strength = None
            if not strength_text:
                row_errors.append("AllowedStrength is required")
            else:
                try:
                    strength = int(strength_text)
                except ValueError:
                    row_errors.append("AllowedStrength must be a whole number")
                if strength is not None and strength < 0:
                    row_errors.append("AllowedStrength cannot be negative")

            if row_errors:
                errors.append(f"Row {offset}: {'; '.join(row_errors)}")
                continue

            seen_ids.add(values["employee_id"])
            values["allowed_strength"] = strength
            parsed.append(values)

        if errors:
            preview = " ".join(errors[:5])
            extra = f" {len(errors) - 5} more." if len(errors) > 5 else ""
            raise HTTPException(status_code=400, detail=preview + extra)
        if not parsed:
            raise HTTPException(status_code=400, detail="The worksheet has no employee rows")

        inserted = 0
        updated = 0
        for values in parsed:
            employee = db.get(Employee, values["employee_id"])
            if employee is None:
                db.add(Employee(**values))
                inserted += 1
            else:
                employee.employee_name = values["employee_name"]
                employee.department = values["department"]
                employee.mobile = values["mobile"]
                employee.company_code = values["company_code"]
                employee.allowed_strength = values["allowed_strength"]
                updated += 1
        db.commit()
        return EmployeeUploadResult(inserted=inserted, updated=updated)
    finally:
        workbook.close()
